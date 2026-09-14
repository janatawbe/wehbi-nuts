"""Tests for Milestone 8 catalog import/export -- Excel (.xlsx) only.

Never touches any AI-dependent endpoint or service -- every Product/
Category here is created directly via the ORM and every import file is
built in-memory, so this whole suite makes zero AI calls and needs no
OpenRouter dependency override at all.
"""
import io
import uuid
from decimal import Decimal

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models.category import Category
from app.models.digitized_product import DigitizedProduct
from app.models.digitization_job import DigitizationJob
from app.models.enums import ReviewStatus, SellingMode, StockStatus
from app.models.product import Product


def make_category(db_session: Session, **overrides) -> Category:
    defaults = dict(name_en="Nuts", name_ar="مكسرات", slug=f"cat-{uuid.uuid4().hex[:8]}")
    defaults.update(overrides)
    category = Category(**defaults)
    db_session.add(category)
    db_session.commit()
    return category


def make_product(db_session: Session, **overrides) -> Product:
    defaults = dict(sku=f"SKU-{uuid.uuid4().hex[:8]}", name_en="Roasted Almonds", name_ar="لوز محمص", price=Decimal("18.00"))
    defaults.update(overrides)
    product = Product(**defaults)
    db_session.add(product)
    db_session.commit()
    return product


def xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


BASE_HEADERS = [
    "ID",
    "SKU",
    "Barcode",
    "English Name",
    "Arabic Name",
    "English Description",
    "Arabic Description",
    "Category",
    "Brand",
    "Selling Mode",
    "Package Weight (kg)",
    "Price",
    "Discount Price",
    "Stock Status",
    "Image",
]


def base_row(**overrides) -> list:
    row = {
        "ID": "",
        "SKU": f"SKU-{uuid.uuid4().hex[:8]}",
        "Barcode": "",
        "English Name": "Roasted Almonds",
        "Arabic Name": "لوز محمص",
        "English Description": "",
        "Arabic Description": "",
        "Category": "",
        "Brand": "",
        "Selling Mode": "",
        "Package Weight (kg)": "",
        "Price": "20.00",
        "Discount Price": "",
        "Stock Status": "",
        "Image": "",
    }
    row.update(overrides)
    return [row[header] for header in BASE_HEADERS]


def upload(client: TestClient, path: str, filename: str, content: bytes) -> "httpx.Response":  # noqa: F821
    return client.post(path, files={"file": (filename, content)})


# --- Export ---------------------------------------------------------------


def test_export_xlsx_returns_a_downloadable_workbook_with_expected_headers(client: TestClient, db_session: Session):
    category = make_category(db_session, name_en="Nuts")
    make_product(db_session, category_id=category.id, price=Decimal("18.00"))

    response = client.get("/api/catalog/export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    header_row = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header_row == BASE_HEADERS


def test_export_xlsx_includes_correct_field_values(client: TestClient, db_session: Session):
    category = make_category(db_session, name_en="Coffee")
    make_product(
        db_session,
        sku="COF-1",
        name_en="Turkish Coffee",
        name_ar="قهوة تركية",
        category_id=category.id,
        brand="Wehbi",
        unit="unit",
        price=Decimal("12.50"),
        stock_status=StockStatus.LOW_STOCK,
    )

    response = client.get("/api/catalog/export.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    row = [cell.value for cell in list(sheet.iter_rows(min_row=2, max_row=2))[0]]
    values = dict(zip(BASE_HEADERS, row))

    assert values["SKU"] == "COF-1"
    assert values["English Name"] == "Turkish Coffee"
    assert values["Arabic Name"] == "قهوة تركية"
    assert values["Category"] == "Coffee"
    assert values["Brand"] == "Wehbi"
    assert values["Selling Mode"] == "unit"
    assert values["Price"] == 12.5
    assert values["Stock Status"] == "low_stock"


def test_export_xlsx_on_an_empty_catalog_still_returns_a_header_only_workbook(client: TestClient):
    response = client.get("/api/catalog/export.xlsx")

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    assert sheet.max_row == 1


def test_export_xlsx_preserves_arabic_unicode_values(client: TestClient, db_session: Session):
    make_product(db_session, sku="AR-1", name_en="Almonds", name_ar="لوز", price=Decimal("9.99"))

    response = client.get("/api/catalog/export.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    row = [cell.value for cell in list(sheet.iter_rows(min_row=2, max_row=2))[0]]
    values = dict(zip(BASE_HEADERS, row))

    assert values["Arabic Name"] == "لوز"
    assert values["Price"] == 9.99


def test_export_reflects_weight_mode_products_correctly(client: TestClient, db_session: Session):
    make_product(
        db_session,
        sku="W-1",
        name_en="Loose Cashews",
        name_ar="كاجو",
        unit="weight",
        weight=None,
        price=Decimal("30.00"),
    )

    response = client.get("/api/catalog/export.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    row = [cell.value for cell in list(sheet.iter_rows(min_row=2, max_row=2))[0]]
    values = dict(zip(BASE_HEADERS, row))

    assert values["Selling Mode"] == "weight"
    assert values["Package Weight (kg)"] is None


def test_export_products_not_linked_to_any_digitized_product_are_still_exported(
    client: TestClient, db_session: Session
):
    """Manually-created Products (no DigitizedProduct behind them at all)
    must export exactly like an M7-approved one -- export never depends
    on DigitizedProduct existing."""
    make_product(db_session, sku="MANUAL-1", name_en="Hand Added", name_ar="مضاف يدويا", price=Decimal("5.00"))

    response = client.get("/api/catalog/export.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    skus = [dict(zip(BASE_HEADERS, [cell.value for cell in row]))["SKU"] for row in sheet.iter_rows(min_row=2)]
    assert "MANUAL-1" in skus


def test_xlsx_export_neutralizes_formula_like_values(client: TestClient, db_session: Session):
    make_product(db_session, sku="F-1", name_en="=cmd|'/c calc'!A1", name_ar="لوز", brand="+SUM(1)")

    response = client.get("/api/catalog/export.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    row = [cell.value for cell in list(sheet.iter_rows(min_row=2, max_row=2))[0]]
    values = dict(zip(BASE_HEADERS, row))

    assert values["English Name"].startswith("'=")
    assert values["Brand"].startswith("'+")


# --- Import: preview never writes -----------------------------------------


def test_preview_makes_zero_database_changes(client: TestClient, db_session: Session):
    category = make_category(db_session, name_en="Nuts")
    content = xlsx_bytes(BASE_HEADERS, [base_row(**{"Category": "Nuts"})])

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)

    assert response.status_code == 200
    assert db_session.query(Product).count() == 0


def test_preview_classifies_a_valid_new_product(client: TestClient, db_session: Session):
    category = make_category(db_session, name_en="Nuts")
    content = xlsx_bytes(BASE_HEADERS, [base_row(**{"Category": "Nuts", "SKU": "NEW-1"})])

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()

    assert body["new_count"] == 1
    assert body["update_count"] == 0
    assert body["invalid_count"] == 0
    assert body["rows"][0]["action"] == "new"
    assert body["rows"][0]["sku"] == "NEW-1"


def test_preview_classifies_an_unchanged_product(client: TestClient, db_session: Session):
    category = make_category(db_session, name_en="Nuts")
    product = make_product(db_session, sku="SAME-1", category_id=category.id, price=Decimal("20.00"))
    content = xlsx_bytes(
        BASE_HEADERS,
        [
            base_row(
                SKU="SAME-1",
                **{
                    "Category": "Nuts",
                    "Price": "20.00",
                    "English Name": product.name_en,
                    "Arabic Name": product.name_ar,
                },
            )
        ],
    )

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()

    assert body["unchanged_count"] == 1
    assert body["rows"][0]["action"] == "unchanged"
    assert body["rows"][0]["changes"] == []


def test_preview_classifies_an_update_and_shows_the_field_diff(client: TestClient, db_session: Session):
    product = make_product(db_session, sku="UPD-1", name_en="Almonds", name_ar="لوز", price=Decimal("18.00"))
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(SKU="UPD-1", **{"English Name": "Almonds", "Arabic Name": "لوز", "Price": "20.00"})],
    )

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()

    assert body["update_count"] == 1
    row = body["rows"][0]
    assert row["action"] == "update"
    price_change = next(c for c in row["changes"] if c["field"] == "price")
    assert price_change["old"] == "18.00"
    assert price_change["new"] == "20.00"


def test_preview_matches_by_product_id_when_present(client: TestClient, db_session: Session):
    product = make_product(db_session, sku="OLD-SKU", name_en="Almonds", name_ar="لوز", price=Decimal("18.00"))
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(ID=str(product.id), SKU="NEW-SKU-RENAMED", **{"English Name": "Almonds", "Arabic Name": "لوز", "Price": "18.00"})],
    )

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()
    row = body["rows"][0]

    assert row["action"] == "update"
    assert row["product_id"] == str(product.id)
    sku_change = next(c for c in row["changes"] if c["field"] == "sku")
    assert sku_change["old"] == "OLD-SKU"
    assert sku_change["new"] == "NEW-SKU-RENAMED"


def test_preview_does_not_match_by_name_alone(client: TestClient, db_session: Session):
    """A row whose name matches an existing product but whose SKU/ID do
    not must be treated as a brand NEW product, never silently matched by
    name (names are editable and unsafe as an identity key)."""
    make_product(db_session, sku="EXISTING-SKU", name_en="Almonds", name_ar="لوز", price=Decimal("18.00"))
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(SKU="DIFFERENT-SKU", **{"English Name": "Almonds", "Arabic Name": "لوز", "Price": "18.00"})],
    )

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()

    assert body["rows"][0]["action"] == "new"
    assert db_session.query(Product).count() == 1


# --- Import: validation -----------------------------------------------------


def test_invalid_negative_price_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row(Price="-5.00")])
    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    row = response.json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("negative" in e for e in row["errors"])


def test_invalid_malformed_price_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row(Price="not-a-number")])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("valid number" in e for e in row["errors"])


def test_missing_price_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row(Price="")])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("Price is required" in e for e in row["errors"])


def test_unknown_category_is_rejected_with_a_clear_message(client: TestClient, db_session: Session):
    content = xlsx_bytes(BASE_HEADERS, [base_row(**{"Category": "Nonexistent Category"})])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("Nonexistent Category" in e for e in row["errors"])
    assert db_session.query(Category).count() == 0


def test_category_never_auto_created(client: TestClient, db_session: Session):
    content = xlsx_bytes(BASE_HEADERS, [base_row(**{"Category": "Brand New Category"})])
    upload(client, "/api/catalog/import/confirm", "products.xlsx", content)
    assert db_session.query(Category).count() == 0


def test_invalid_selling_mode_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row(**{"Selling Mode": "per-kg"})])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("Selling mode" in e for e in row["errors"])


def test_weight_mode_with_package_weight_is_rejected(client: TestClient):
    content = xlsx_bytes(
        BASE_HEADERS, [base_row(**{"Selling Mode": "weight", "Package Weight (kg)": "0.5"})]
    )
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("package weight" in e for e in row["errors"])


def test_missing_required_name_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row(**{"English Name": ""})])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("English name" in e for e in row["errors"])


def test_duplicate_sku_within_the_same_file_is_flagged_and_not_applied(client: TestClient, db_session: Session):
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(SKU="DUP-SKU"), base_row(SKU="DUP-SKU")],
    )
    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()
    assert all(row["action"] == "invalid" for row in body["rows"])
    assert all(any("Duplicate SKU" in e for e in row["errors"]) for row in body["rows"])

    confirm = upload(client, "/api/catalog/import/confirm", "products.xlsx", content)
    assert confirm.json()["created"] == 0
    assert db_session.query(Product).count() == 0


def test_duplicate_product_id_within_the_same_file_is_flagged_and_not_applied(
    client: TestClient, db_session: Session
):
    product = make_product(db_session, sku="EXISTING")
    content = xlsx_bytes(
        BASE_HEADERS,
        [
            base_row(ID=str(product.id), SKU="ROW-A", Price="10.00"),
            base_row(ID=str(product.id), SKU="ROW-B", Price="11.00"),
        ],
    )
    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()
    assert all(row["action"] == "invalid" for row in body["rows"])
    assert all(any("Duplicate Product ID" in e for e in row["errors"]) for row in body["rows"])


def test_product_id_that_does_not_exist_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row(ID=str(uuid.uuid4()))])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("does not match any existing product" in e for e in row["errors"])


def test_malformed_product_id_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row(ID="not-a-uuid")])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("not a valid Product ID" in e for e in row["errors"])


def test_sku_clashing_with_a_different_existing_product_is_rejected(client: TestClient, db_session: Session):
    other = make_product(db_session, sku="TAKEN-SKU")
    target = make_product(db_session, sku="TARGET-SKU")
    content = xlsx_bytes(BASE_HEADERS, [base_row(ID=str(target.id), SKU="TAKEN-SKU")])
    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "invalid"
    assert any("already used by another product" in e for e in row["errors"])


def test_missing_required_columns_is_rejected_at_the_file_level(client: TestClient):
    headers = [h for h in BASE_HEADERS if h != "Price"]
    content = xlsx_bytes(headers, [[v for h, v in zip(BASE_HEADERS, base_row()) if h != "Price"]])
    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    assert response.status_code == 400
    assert "Price" in response.json()["detail"]


def test_duplicate_headers_in_the_file_is_rejected(client: TestClient):
    headers = BASE_HEADERS + ["SKU"]
    row = base_row() + ["extra"]
    content = xlsx_bytes(headers, [row])
    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    assert response.status_code == 400
    assert "Duplicate column header" in response.json()["detail"]


def test_wrong_file_type_is_rejected(client: TestClient):
    response = upload(client, "/api/catalog/import/preview", "products.txt", b"not a spreadsheet")
    assert response.status_code == 400
    assert "Please upload an Excel (.xlsx) file." in response.json()["detail"]


def test_csv_file_is_no_longer_accepted(client: TestClient):
    content = "SKU,English Name\nSKU-1,Almonds\n".encode("utf-8-sig")
    response = upload(client, "/api/catalog/import/preview", "products.csv", content)
    assert response.status_code == 400
    assert "Please upload an Excel (.xlsx) file." in response.json()["detail"]


def test_legacy_xls_file_is_rejected(client: TestClient):
    content = xlsx_bytes(BASE_HEADERS, [base_row()])
    response = upload(client, "/api/catalog/import/preview", "products.xls", content)
    assert response.status_code == 400
    assert "Please upload an Excel (.xlsx) file." in response.json()["detail"]


def test_malformed_xlsx_content_is_rejected(client: TestClient):
    response = upload(client, "/api/catalog/import/preview", "products.xlsx", b"this is not a real workbook")
    assert response.status_code == 400
    assert "valid .xlsx" in response.json()["detail"]


def test_empty_file_is_rejected(client: TestClient):
    response = upload(client, "/api/catalog/import/preview", "products.xlsx", b"")
    assert response.status_code == 400


def test_file_too_large_is_rejected(client: TestClient):
    from app.core.config import Settings, get_settings
    from app.main import app

    content = xlsx_bytes(BASE_HEADERS, [base_row()])
    app.dependency_overrides[get_settings] = lambda: Settings(catalog_import_max_file_size_bytes=len(content) - 1)

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)

    assert response.status_code == 400
    assert "too large" in response.json()["detail"]


# --- Import: Arabic/Unicode -------------------------------------------------


def test_valid_arabic_row_is_previewed_as_new(client: TestClient, db_session: Session):
    category = make_category(db_session, name_en="Nuts")
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(**{"Category": "Nuts", "Arabic Name": "لوز مملح", "English Name": "Salted Almonds"})],
    )

    response = upload(client, "/api/catalog/import/preview", "products.xlsx", content)
    body = response.json()

    assert body["new_count"] == 1
    assert body["rows"][0]["name_en"] == "Salted Almonds"


def test_category_resolves_case_insensitively(client: TestClient, db_session: Session):
    make_category(db_session, name_en="Nuts", slug="nuts")
    content = xlsx_bytes(BASE_HEADERS, [base_row(**{"Category": "nuts"})])

    row = upload(client, "/api/catalog/import/preview", "products.xlsx", content).json()["rows"][0]
    assert row["action"] == "new"
    assert row["errors"] == []


# --- Import: confirm applies changes ----------------------------------------


def test_confirm_creates_a_valid_new_product(client: TestClient, db_session: Session):
    category = make_category(db_session, name_en="Nuts")
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(SKU="CREATE-1", **{"Category": "Nuts", "Price": "15.00"})],
    )

    response = upload(client, "/api/catalog/import/confirm", "products.xlsx", content)
    body = response.json()

    assert body["created"] == 1
    assert body["updated"] == 0
    assert body["failed"] == 0
    created = db_session.query(Product).filter_by(sku="CREATE-1").one()
    assert created.price == Decimal("15.00")
    assert created.category_id == category.id


def test_confirm_updates_an_existing_product_in_place_preserving_its_id(client: TestClient, db_session: Session):
    product = make_product(db_session, sku="UPD-2", name_en="Almonds", name_ar="لوز", price=Decimal("18.00"))
    original_id = product.id
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(ID=str(product.id), SKU="UPD-2", **{"English Name": "Almonds", "Arabic Name": "لوز", "Price": "22.00"})],
    )

    response = upload(client, "/api/catalog/import/confirm", "products.xlsx", content)
    body = response.json()

    assert body["updated"] == 1
    assert body["created"] == 0
    assert db_session.query(Product).count() == 1
    refreshed = db_session.get(Product, original_id)
    assert refreshed.id == original_id
    assert refreshed.price == Decimal("22.00")


def test_confirm_skips_unchanged_rows(client: TestClient, db_session: Session):
    product = make_product(db_session, sku="SAME-2", name_en="Almonds", name_ar="لوز", price=Decimal("20.00"))
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(SKU="SAME-2", **{"English Name": "Almonds", "Arabic Name": "لوز", "Price": "20.00"})],
    )

    response = upload(client, "/api/catalog/import/confirm", "products.xlsx", content)
    body = response.json()

    assert body["unchanged"] == 1
    assert body["created"] == 0
    assert body["updated"] == 0


def test_confirm_never_applies_invalid_rows_and_reports_them_explicitly(client: TestClient, db_session: Session):
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(SKU="GOOD-1", Price="10.00"), base_row(SKU="BAD-1", Price="-10.00")],
    )

    response = upload(client, "/api/catalog/import/confirm", "products.xlsx", content)
    body = response.json()

    assert body["created"] == 1
    assert body["failed"] == 1
    assert body["failures"][0]["sku"] == "BAD-1"
    assert any("negative" in e for e in body["failures"][0]["errors"])
    assert db_session.query(Product).count() == 1


def test_confirm_does_not_duplicate_an_existing_product(client: TestClient, db_session: Session):
    product = make_product(db_session, sku="NODUP-1", name_en="Almonds", name_ar="لوز", price=Decimal("18.00"))
    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(SKU="NODUP-1", **{"English Name": "Almonds", "Arabic Name": "لوز", "Price": "19.00"})],
    )

    upload(client, "/api/catalog/import/confirm", "products.xlsx", content)

    assert db_session.query(Product).filter_by(sku="NODUP-1").count() == 1


# --- Import: Milestone 7 linkage protection ---------------------------------


def test_m7_linked_product_updated_by_import_keeps_the_same_id_and_digitized_product_linkage(
    client: TestClient, db_session: Session
):
    """The exact scenario Part 7 is about: a spreadsheet edits a Product
    that an approved DigitizedProduct already points to. The Product row
    must be updated in place; the DigitizedProduct's own review record and
    its product_id linkage must be completely untouched."""
    job = DigitizationJob()
    db_session.add(job)
    db_session.commit()

    category = make_category(db_session, name_en="Nuts")
    product = make_product(db_session, sku="M7-LINKED", name_en="Almonds", name_ar="لوز", price=Decimal("18.00"), category_id=category.id)
    digitized = DigitizedProduct(
        job_id=job.id,
        name_en="Almonds",
        name_ar="لوز",
        category_id=category.id,
        selling_mode=SellingMode.UNIT,
        price=Decimal("18.00"),
        review_status=ReviewStatus.APPROVED,
        product_id=product.id,
    )
    db_session.add(digitized)
    db_session.commit()
    digitized_id = digitized.id

    content = xlsx_bytes(
        BASE_HEADERS,
        [base_row(ID=str(product.id), SKU="M7-LINKED", **{"Category": "Nuts", "English Name": "Almonds", "Arabic Name": "لوز", "Price": "25.00"})],
    )
    response = upload(client, "/api/catalog/import/confirm", "products.xlsx", content)
    assert response.json()["updated"] == 1

    refreshed_product = db_session.get(Product, product.id)
    assert refreshed_product.price == Decimal("25.00")

    refreshed_digitized = db_session.get(DigitizedProduct, digitized_id)
    assert refreshed_digitized.product_id == product.id
    assert refreshed_digitized.review_status == ReviewStatus.APPROVED
    # The historical review record's own fields are untouched by a catalog
    # spreadsheet edit -- Product and DigitizedProduct stay separate
    # concepts (see Part 7).
    assert refreshed_digitized.price == Decimal("18.00")


def test_import_never_creates_a_second_product_for_an_m7_approved_product(client: TestClient, db_session: Session):
    job = DigitizationJob()
    db_session.add(job)
    db_session.commit()
    product = make_product(db_session, sku="M7-2", name_en="Cashews", name_ar="كاجو", price=Decimal("30.00"))
    digitized = DigitizedProduct(
        job_id=job.id,
        name_en="Cashews",
        name_ar="كاجو",
        selling_mode=SellingMode.UNIT,
        price=Decimal("30.00"),
        review_status=ReviewStatus.APPROVED,
        product_id=product.id,
    )
    db_session.add(digitized)
    db_session.commit()

    content = xlsx_bytes(BASE_HEADERS, [base_row(ID=str(product.id), SKU="M7-2", **{"English Name": "Cashews", "Arabic Name": "كاجو", "Price": "31.00"})])
    upload(client, "/api/catalog/import/confirm", "products.xlsx", content)

    assert db_session.query(Product).count() == 1
